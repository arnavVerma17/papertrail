"""Text extraction and source-aware chunking for common research documents."""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re

import fitz


@dataclass(frozen=True)
class TextChunk:
    text: str
    source: str
    page: int
    location: str | None = None

    @property
    def citation_location(self) -> str:
        return self.location or f"Page {self.page}"


_QUESTION_HEADING = re.compile(
    r"^\s*\d{1,3}[.)]\s+(?:explain|define|discuss|what|how|compare|differentiate|"
    r"describe|list|state|write|give|illustrate|analyze|analyse|identify|"
    r"distinguish|elaborate|show|mention|outline|why|when|where|which|who|"
    r"summari[sz]e|examine|evaluate|demonstrate)\b",
    re.IGNORECASE,
)


def _page_sections(text: str) -> list[str]:
    """Keep numbered question-and-answer sections together when a PDF has them."""
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines()]
    lines = [line for line in lines if line]
    if not lines:
        return []

    headings = [index for index, line in enumerate(lines) if _QUESTION_HEADING.match(line)]
    if not headings:
        return [" ".join(lines)]

    starts = ([0] if headings[0] > 0 else []) + headings
    sections = []
    for position, start in enumerate(starts):
        end = starts[position + 1] if position + 1 < len(starts) else len(lines)
        section = " ".join(lines[start:end]).strip()
        if section:
            sections.append(section)
    return sections


def extract_chunks(
    pdf_bytes: bytes,
    filename: str,
    chunk_words: int = 180,
    overlap: int = 35,
    ocr: bool = False,
    ocr_language: str = "eng",
) -> list[TextChunk]:
    """Extract page text, optionally OCRing pages with little selectable text."""
    chunks: list[TextChunk] = []
    with fitz.open(stream=pdf_bytes, filetype="pdf") as document:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if ocr and len(text) < 60:
                text_page = page.get_textpage_ocr(
                    language=ocr_language,
                    dpi=200,
                    full=not bool(text),
                )
                text = page.get_text("text", textpage=text_page).strip()
            step = max(1, chunk_words - overlap)
            for section in _page_sections(text):
                words = section.split()
                for start in range(0, len(words), step):
                    piece = " ".join(words[start : start + chunk_words]).strip()
                    if piece:
                        chunks.append(TextChunk(piece, filename, page_number))
                    if start + chunk_words >= len(words):
                        break
    return chunks


def _numbered_sections(
    units: list[tuple[int, str]], styled_headings: set[int] | None = None
) -> list[list[tuple[int, str]]]:
    """Split numbered question-and-answer documents at detected section headings."""
    styled_headings = styled_headings or set()
    headings = [
        index
        for index, (_, text) in enumerate(units)
        if index in styled_headings or _QUESTION_HEADING.match(text)
    ]
    if not headings:
        return [units] if units else []
    starts = ([0] if headings[0] > 0 else []) + headings
    return [units[start : starts[index + 1] if index + 1 < len(starts) else len(units)]
            for index, start in enumerate(starts)]


def _chunk_units(
    units: list[tuple[int, str]],
    filename: str,
    chunk_words: int,
    overlap: int,
    location_kind: str,
    next_location_id: int = 1,
    fixed_location: tuple[int, str] | None = None,
) -> tuple[list[TextChunk], int]:
    """Chunk text while retaining the source line or paragraph range."""
    tokens = [
        (word, unit_number)
        for unit_number, text in units
        for word in text.split()
    ]
    if not tokens:
        return [], next_location_id

    chunks = []
    step = max(1, chunk_words - overlap)
    for start in range(0, len(tokens), step):
        window = tokens[start : start + chunk_words]
        content = " ".join(word for word, _ in window).strip()
        if fixed_location is None:
            location_id = next_location_id
            location = f"{location_kind} {window[0][1]}–{window[-1][1]}"
            next_location_id += 1
        else:
            location_id, location = fixed_location
        if content:
            chunks.append(TextChunk(content, filename, location_id, location))
        if start + chunk_words >= len(tokens):
            break
    return chunks, next_location_id


def _extract_docx_chunks(
    file_bytes: bytes, filename: str, chunk_words: int, overlap: int
) -> list[TextChunk]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(BytesIO(file_bytes))
    blocks: list[tuple[int, str]] = []
    styled_headings: set[int] = set()
    for element in document.element.body.iterchildren():
        is_heading = False
        if element.tag.endswith("}p"):
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            style_name = paragraph.style.name.casefold() if paragraph.style else ""
            is_heading = style_name.startswith("heading") or style_name in {"title", "subtitle"}
        elif element.tag.endswith("}tbl"):
            table = Table(element, document)
            text = "\n".join(
                " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                for row in table.rows
            ).strip()
        else:
            continue
        if text:
            if is_heading:
                styled_headings.add(len(blocks))
            blocks.append((len(blocks) + 1, text))

    chunks: list[TextChunk] = []
    next_location_id = 1
    for section in _numbered_sections(blocks, styled_headings):
        section_chunks, next_location_id = _chunk_units(
            section, filename, chunk_words, overlap, "Content blocks", next_location_id
        )
        chunks.extend(section_chunks)
    return chunks


def _extract_pptx_chunks(
    file_bytes: bytes, filename: str, chunk_words: int, overlap: int
) -> list[TextChunk]:
    from pptx import Presentation

    presentation = Presentation(BytesIO(file_bytes))
    chunks: list[TextChunk] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        units = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                units.append((len(units) + 1, shape.text.strip()))
            if getattr(shape, "has_table", False):
                table_text = "\n".join(
                    " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    for row in shape.table.rows
                ).strip()
                if table_text:
                    units.append((len(units) + 1, table_text))
        slide_chunks, _ = _chunk_units(
            units,
            filename,
            chunk_words,
            overlap,
            "",
            fixed_location=(slide_number, f"Slide {slide_number}"),
        )
        chunks.extend(slide_chunks)
    return chunks


def _extract_text_chunks(
    file_bytes: bytes, filename: str, chunk_words: int, overlap: int
) -> list[TextChunk]:
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = file_bytes.decode("cp1252", errors="replace")
    lines = [(line_number, line.strip()) for line_number, line in enumerate(text.splitlines(), start=1) if line.strip()]
    chunks, _ = _chunk_units(lines, filename, chunk_words, overlap, "Lines")
    return chunks


def extract_document_chunks(
    file_bytes: bytes,
    filename: str,
    chunk_words: int = 180,
    overlap: int = 35,
    ocr: bool = False,
    ocr_language: str = "eng",
) -> list[TextChunk]:
    """Extract searchable text from PDF, DOCX, PPTX, TXT, or Markdown files."""
    extension = Path(filename).suffix.casefold()
    if extension == ".pdf":
        return extract_chunks(file_bytes, filename, chunk_words, overlap, ocr, ocr_language)
    if extension == ".docx":
        return _extract_docx_chunks(file_bytes, filename, chunk_words, overlap)
    if extension == ".pptx":
        return _extract_pptx_chunks(file_bytes, filename, chunk_words, overlap)
    if extension in {".txt", ".md"}:
        return _extract_text_chunks(file_bytes, filename, chunk_words, overlap)
    raise ValueError("Supported formats are PDF, DOCX, PPTX, TXT, and MD.")
